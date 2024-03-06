from sympy import symbols, simplify, Function, solve, MatrixSymbol
from numpy import arange, array
from numpy import linalg as LA
from docx import Document
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

doc = Document()
table = doc.add_table(rows=1, cols=6)

p1 = 8.4e-6
p2 = 6.6667e-4
p3 = 1.7778e-5
p5 = 2
p6_arr = [10, 100]
jacobian_matrix = []
jacobian_matrix_arr = []
x1_solve = []
x2_solve = []
x3_solve = []
p_step = [0.01, 0.1, 1, 10, 100, 1000]
for p6 in p6_arr:  # разные значения p6
    print("p6 = ", p6)
    j_temp = []
    j_temp_prev = []
    # for step in range(0, len(p_step)-1):
    for step in range(0, len(p_step)-1):
        b = p_step[step]
        e = p_step[step+1]
        s = b
        print("interval ", b, "; ", e)
        cells = table.add_row().cells
        for i in range(4):
            cells[i].text = str("p4["+str(b)+';'+str(e)+')')
        for p4 in arange(b, e, s):  # диапазон по р4
            # x1 = symbols('x1', real=True)
            x1 = symbols('x1')
            f = Function('f')
            f = (-p1*(x1**2-x1+p4*x1*p2)/(p1-x1) - (x1**3-x1**2+p4*x1*x1*p2) /
                 (p1-x1) + p5*x1/(1 + p4))/p3 + p4*(p6 - (x1**2-x1+p4*x1*p2)/(p1-x1))
            x1_solution = array(solve(f, x1), dtype=complex)
            # for ch in x1_solution:
            #     if (ch <= complex(0, 0)):
            #         print("skip: x1 <= 0")
            #         continue
            x1_solve_temp = []
            # x1_solve.append(x1_solve_temp)
            # print("x1: ", round(x1_solution[0], 20), round(
            #     x1_solution[1], 20), round(x1_solution[2], 20))
            x1_solve_temp = [x1_solution[0].real,
                             x1_solution[1].real, x1_solution[2].real]
            # if (x1_solve_temp[0] < 0 or x1_solve_temp[1] < 0 or x1_solve_temp[2] < 0):
            #     print("skip: x1 <= 0")
            #     continue
            x1_solve.append(x1_solve_temp)
            x2_solve_temp = []
            x3_solve_temp = []

            for x1_temp in x1_solve_temp:
                x2_solve_temp.append(
                    (x1_temp**2-x1_temp+p4*x1_temp*p2)/(p1-x1_temp))
                x3_solve_temp.append((x1_temp/(1 + p4)))
            x2_solve.append(x2_solve_temp)
            x3_solve.append(x3_solve_temp)
            for i in range(3):
                curr_st = ""
                past_st = "stable"
                cells = table.add_row().cells
                table.style = 'Table Grid'
                cells[0].text = str("p4="+str(p4))
                if (x1_solve_temp[i] <= 0 or x2_solve_temp[i] <= 0 or x3_solve_temp[i] <= 0):
                    print("skip: <= 0")
                    for i in range(1, 4):
                        cells[i].text = "skip"
                    jacobian_matrix_arr.append([0, 0, 0, "none"])
                    continue
                else:
                    jacobian_matrix = array([[(-x2_solve_temp[i]+1-2*x1_solve_temp[i])/p2 - p4, (p1-x1_solve_temp[i])/p2, 0],
                                            [-x2_solve_temp[i]/p3,
                                            (-p1-x1_solve_temp[i])/p3-p4, p5/p3],
                                            [1, 0, -1-p4]], dtype=complex)
                    # for pr in jacobian_matrix:
                    #     print(pr)
                    wa, va = LA.eig(jacobian_matrix)
                    # for i in range(1, 4):
                    #     cells[i].text = str(wa[i])
                    if (wa[0].real < 0 and wa[1].real < 0 and wa[2].real < 0):
                        j_temp = [wa[0], wa[1], wa[2], "stable"]
                    elif (wa[0].real > 0 or wa[1].real > 0 or wa[2].real > 0):
                        j_temp = [wa[0], wa[1], wa[2], "unstable"]
                    else:
                        j_temp = [wa[0], wa[1], wa[2], "check"]
                    print(j_temp)
                    jacobian_matrix_arr.append(j_temp)
                    curr_st = j_temp[3]
                    # добавить вывод шага
                    if ((curr_st == "stable" and past_st == "unstable") or (curr_st == "unstable" and past_st == "stable")):
                        for i in range(1, 5):
                            shading_elm_1 = parse_xml(
                                r'<w:shd {} w:fill="90EE90"/>'.format(nsdecls('w')))
                            cells[i]._tc.get_or_add_tcPr().append(
                                shading_elm_1)
                            cells[i].text = str(j_temp[i-1])
                    else:
                        for i in range(1, 5):
                            cells[i].text = str(j_temp[i-1])
                    past_st = curr_st
            cells = table.add_row().cells
            for i in range(4):
                cells[i].text = "---"
            print("next step")

            # for i in range(9):
            #     for j in range(3):
            #         cell = table.cell(i, j)
            #         if (x3_solve[i][j] <= 0):
            #             cell.text = '-'
            #         else:
            #             cell.text = str(x3_solve[i][j])

            # table.add_row()
            # wa, va = LA.eig(jacobian_matrix)
            # print(wa)
            # LA.det(jacobian_matrix)

            # LA.det(jacobian_matrix)

            # print(jacobian_matrix)
            # print(x1_solve)
            # print()
        print()
doc.save('jacobian_st1.doc')
